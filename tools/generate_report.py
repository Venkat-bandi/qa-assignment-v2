"""
Auto-generates a full test-case report with:
  - TC ID, Description, Pre-conditions, Test Data, Steps
  - Expected Result, Actual Result, Status

Reads:   reports/behave.json
Writes:  reports/test_report.xlsx
         reports/test_report.html
         reports/bugs.md
"""
import json
import os
from datetime import datetime


JSON_PATH = "reports/behave.json"
XLSX_PATH = "reports/test_report.xlsx"
HTML_PATH = "reports/test_report.html"
BUGS_PATH = "reports/bugs.md"


SCENARIO_META = {
    "Register a new user and log in successfully": {
        "id": "TC-01",
        "desc": "Verify new user can register and then log in",
        "pre": "None",
        "data": "Unique email; Password: Test@123",
        "expected": "Redirect to /login, then /products after login",
    },
    "Login fails with wrong password": {
        "id": "TC-02",
        "desc": "Verify login fails for invalid credentials",
        "pre": "User is registered",
        "data": "Email: testuser@shop.com; Password: WrongPass123",
        "expected": "Error message displayed; user stays on login page",
    },
    "Add a product to the cart and verify count": {
        "id": "TC-03",
        "desc": "Verify user can add a product and cart count updates",
        "pre": "User logged in",
        "data": "Product: iPhone; Qty: 1",
        "expected": "Cart contains 1 item; 'iPhone' displayed",
    },
    "Add multiple different products to the cart": {
        "id": "TC-04",
        "desc": "Verify multiple products can be added",
        "pre": "User logged in",
        "data": "iPhone x1; MacBook x1",
        "expected": "Cart contains 2 items",
    },
    "Update quantity of an existing product in cart": {
        "id": "TC-05",
        "desc": "Verify quantity can be set on a product",
        "pre": "User logged in",
        "data": "iPad; Qty: 2",
        "expected": "Cart has 1 row with quantity = 2",
    },
    "Apply a valid coupon": {
        "id": "TC-06",
        "desc": "Verify valid coupon applies 10% discount",
        "pre": "User logged in; product in cart",
        "data": "Coupon: SAVE10",
        "expected": "Coupon accepted; discount applied; total reduced",
    },
    "Apply an expired coupon": {
        "id": "TC-07",
        "desc": "Verify expired coupon is rejected",
        "pre": "User logged in; product in cart",
        "data": "Coupon: EXPIRED2020",
        "expected": "Coupon rejected; discount remains -$0.00",
    },
    "Apply an invalid coupon": {
        "id": "TC-08",
        "desc": "Verify unknown coupon is rejected",
        "pre": "User logged in; product in cart",
        "data": "Coupon: FAKE123",
        "expected": "Coupon rejected; discount remains -$0.00",
    },
    "Place an order successfully": {
        "id": "TC-09",
        "desc": "Verify order can be placed successfully",
        "pre": "User logged in; product in cart",
        "data": "iPhone x1",
        "expected": "Redirect to /orders; at least 1 order visible",
    },
    'App rejects invalid quantity "0"': {
        "id": "TC-10",
        "desc": "Verify app rejects quantity = 0",
        "pre": "User logged in",
        "data": "Quantity: 0",
        "expected": "Item rejected; cart contains 0 items",
    },
    'App rejects invalid quantity "-1"': {
        "id": "TC-11",
        "desc": "Verify app rejects negative quantity",
        "pre": "User logged in",
        "data": "Quantity: -1",
        "expected": "Item rejected; cart contains 0 items",
    },
    'App rejects invalid quantity "9999"': {
        "id": "TC-12",
        "desc": "Verify app rejects quantity above max (10)",
        "pre": "User logged in",
        "data": "Quantity: 9999",
        "expected": "Item rejected; cart contains 0 items",
    },
    "App accepts maximum allowed quantity": {
        "id": "TC-13",
        "desc": "Verify app accepts max allowed quantity",
        "pre": "User logged in",
        "data": "Quantity: 10",
        "expected": "Cart contains 1 item with quantity = 10",
    },
}


def load_json():
    if not os.path.exists(JSON_PATH):
        raise FileNotFoundError(f"{JSON_PATH} not found. Run behave first.")
    with open(JSON_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


def match_meta(scenario_name):
    if scenario_name in SCENARIO_META:
        return SCENARIO_META[scenario_name]

    base = scenario_name.split(" -- ")[0].strip()
    for key, val in SCENARIO_META.items():
        if base.startswith(key) or key.startswith(base):
            return val

    if "invalid quantity" in scenario_name.lower():
        parts = scenario_name.split('"')
        qty = parts[1] if len(parts) > 1 else "?"
        return {
            "id": "—",
            "desc": f"Verify app rejects quantity = {qty}",
            "pre": "User logged in",
            "data": f"Quantity: {qty}",
            "expected": "Item rejected; cart contains 0 items",
        }
    return {}


def extract_scenarios(data):
    rows = []
    for feature in data:
        for element in feature.get("elements", []):
            if element.get("type") != "scenario":
                continue

            sname = element.get("name", "")
            meta = match_meta(sname)

            steps_desc = " → ".join(
                f"{s.get('keyword','').strip()} {s.get('name','')}"
                for s in element.get("steps", [])
            )

            failed_step = ""
            for s in element.get("steps", []):
                if s.get("status") == "failed":
                    failed_step = s.get("name", "")
                    break

            status = element.get("status", "unknown")
            actual = "As expected" if status == "passed" else f"Failed at: {failed_step}"

            rows.append({
                "tc_id": meta.get("id", "—"),
                "feature": feature.get("name", ""),
                "scenario": sname,
                "description": meta.get("desc", sname),
                "pre_conditions": meta.get("pre", "—"),
                "test_data": meta.get("data", "—"),
                "steps": steps_desc,
                "expected": meta.get("expected", "—"),
                "actual": actual,
                "status": status.upper(),
                "duration": round(element.get("duration", 0) / 1e9, 2),
                "failed_step": failed_step,
            })
    return rows


def write_excel(rows):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        print("⚠️  openpyxl not installed. Run: pip install openpyxl")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Test Cases"

    headers = [
        "TC ID", "Feature", "Scenario", "Description",
        "Pre-conditions", "Test Data", "Steps",
        "Expected Result", "Actual Result", "Status", "Duration (s)"
    ]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="4F46E5")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    pass_fill = PatternFill("solid", fgColor="D1FAE5")
    fail_fill = PatternFill("solid", fgColor="FEE2E2")

    for i, r in enumerate(rows, start=2):
        ws.append([
            r["tc_id"], r["feature"], r["scenario"], r["description"],
            r["pre_conditions"], r["test_data"], r["steps"],
            r["expected"], r["actual"], r["status"], r["duration"],
        ])
        status_cell = ws.cell(row=i, column=10)
        status_cell.fill = pass_fill if r["status"] == "PASSED" else fail_fill
        status_cell.alignment = Alignment(horizontal="center")
        status_cell.font = Font(bold=True)

    widths = [8, 20, 34, 34, 20, 22, 55, 38, 34, 12, 12]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + idx)].width = w

    ws.freeze_panes = "A2"
    wb.save(XLSX_PATH)
    print(f"✅ Excel report → {XLSX_PATH}")


def write_html(rows):
    total = len(rows)
    passed = sum(1 for r in rows if r["status"] == "PASSED")
    failed = total - passed
    pct = (passed / total * 100) if total else 0

    trs = ""
    for r in rows:
        cls = "pass" if r["status"] == "PASSED" else "fail"
        icon = "✅" if r["status"] == "PASSED" else "❌"
        trs += f"""
        <tr class="{cls}">
            <td><strong>{r['tc_id']}</strong></td>
            <td>{r['feature']}</td>
            <td>{r['scenario']}</td>
            <td>{r['description']}</td>
            <td>{r['pre_conditions']}</td>
            <td>{r['test_data']}</td>
            <td class="steps">{r['steps']}</td>
            <td>{r['expected']}</td>
            <td>{r['actual']}</td>
            <td>{icon} {r['status']}</td>
            <td>{r['duration']}s</td>
        </tr>"""

    html = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>ShopQA Test Report</title>
<style>
  body {{ font-family: system-ui, sans-serif; background:#f9fafb; margin:0; padding:24px; color:#111827; font-size:12px; }}
  h1 {{ margin: 0 0 4px; }}
  .meta {{ color:#6b7280; margin-bottom:20px; }}
  .cards {{ display:flex; gap:12px; margin-bottom:20px; }}
  .card {{ background:#fff; border-radius:10px; padding:16px 24px; box-shadow:0 1px 3px rgba(0,0,0,0.08); text-align:center; }}
  .card .num {{ font-size:26px; font-weight:700; }}
  .card.pass .num {{ color:#059669; }}
  .card.fail .num {{ color:#dc2626; }}
  table {{ width:100%; border-collapse:collapse; background:#fff; border-radius:10px; overflow:hidden; box-shadow:0 1px 3px rgba(0,0,0,0.08); }}
  th, td {{ padding:10px 12px; text-align:left; border-bottom:1px solid #f3f4f6; vertical-align:top; }}
  th {{ background:#4f46e5; color:#fff; font-weight:600; position:sticky; top:0; }}
  tr.pass td:nth-child(10) {{ color:#059669; font-weight:600; }}
  tr.fail td:nth-child(10) {{ color:#dc2626; font-weight:600; }}
  tr.fail {{ background:#fef2f2; }}
  .steps {{ font-size:11px; color:#6b7280; max-width: 280px; }}
</style></head><body>
  <h1>🧪 ShopQA Automation Test Report</h1>
  <div class="meta">Generated {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}</div>

  <div class="cards">
    <div class="card"><div class="num">{total}</div>Total</div>
    <div class="card pass"><div class="num">{passed}</div>Passed</div>
    <div class="card fail"><div class="num">{failed}</div>Failed</div>
    <div class="card"><div class="num">{pct:.1f}%</div>Pass Rate</div>
  </div>

  <table>
    <thead><tr>
      <th>TC ID</th><th>Feature</th><th>Scenario</th><th>Description</th>
      <th>Pre-conditions</th><th>Test Data</th><th>Steps</th>
      <th>Expected</th><th>Actual</th><th>Status</th><th>Duration</th>
    </tr></thead>
    <tbody>{trs}</tbody>
  </table>
</body></html>"""

    with open(HTML_PATH, "w", encoding="utf-8") as f:
        f.write(html)
    print(f"✅ HTML report → {HTML_PATH}")


def write_bugs(rows):
    failed = [r for r in rows if r["status"] != "PASSED"]
    with open(BUGS_PATH, "w", encoding="utf-8") as f:
        f.write("# Auto-Generated Bug Log\n\n")
        f.write(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")
        if not failed:
            f.write("✅ No defects found in this run.\n")
        else:
            f.write(f"**Total Defects:** {len(failed)}\n\n")
            for i, r in enumerate(failed, start=1):
                f.write(f"## BUG-{i:03d} — {r['scenario']}\n\n")
                f.write(f"| Field | Value |\n|---|---|\n")
                f.write(f"| TC ID | {r['tc_id']} |\n")
                f.write(f"| Feature | {r['feature']} |\n")
                f.write(f"| Description | {r['description']} |\n")
                f.write(f"| Pre-conditions | {r['pre_conditions']} |\n")
                f.write(f"| Test Data | {r['test_data']} |\n")
                f.write(f"| Steps | {r['steps']} |\n")
                f.write(f"| Expected | {r['expected']} |\n")
                f.write(f"| Actual | {r['actual']} |\n")
                f.write(f"| Status | ❌ FAILED |\n\n")
    print(f"✅ Bug log → {BUGS_PATH}")


def main():
    os.makedirs("reports", exist_ok=True)
    data = load_json()
    rows = extract_scenarios(data)

    print(f"\n📊 Generating report for {len(rows)} scenarios...\n")
    write_excel(rows)
    write_html(rows)
    write_bugs(rows)
    print("\n✨ All reports generated in reports/\n")


if __name__ == "__main__":
    main()